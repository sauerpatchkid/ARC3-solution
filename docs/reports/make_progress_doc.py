"""Build docs/reports/Goose_Semester2_Progress.docx (1-2 pages, plain language, for the advisor).

    uv run --with python-docx python docs/reports/make_progress_doc.py docs/reports/Goose_Semester2_Progress.docx
"""
import sys

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

OUT = sys.argv[1]
ACCENT = RGBColor(0x1F, 0x4E, 0x79)

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
for side in ("left_margin", "right_margin"):
    setattr(sec, side, Inches(0.9))
sec.top_margin = sec.bottom_margin = Inches(0.8)

normal = doc.styles["Normal"]
normal.font.name = "Calibri"
normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
normal.font.size = Pt(11)
normal.paragraph_format.space_after = Pt(4)
normal.paragraph_format.line_spacing = 1.08


def heading(text, size=13, before=10):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    r.bold = True
    r.font.size = Pt(size)
    r.font.color.rgb = ACCENT
    return p


def para(text, bold_lead=None, italic=False, size=None, after=None):
    p = doc.add_paragraph()
    if bold_lead:
        p.add_run(bold_lead + " ").bold = True
    r = p.add_run(text)
    r.italic = italic
    if size:
        for run in p.runs:
            run.font.size = Pt(size)
    if after is not None:
        p.paragraph_format.space_after = Pt(after)
    return p


def bullet(text, bold_lead=None, style="List Bullet"):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_after = Pt(3)
    if bold_lead:
        p.add_run(bold_lead + " ").bold = True
    p.add_run(text)
    return p


def shade(cell, hex_fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tcPr.append(shd)


# ---- title ---------------------------------------------------------------
t = doc.add_paragraph()
t.paragraph_format.space_after = Pt(0)
r = t.add_run("Teaching Goose to Remember")
r.bold = True
r.font.size = Pt(18)
r.font.color.rgb = ACCENT
st = doc.add_paragraph()
st.paragraph_format.space_after = Pt(6)
r = st.add_run("ARC-AGI-3 Capstone, Team B  |  Semester 2 progress and next step  |  Matt Sauer, September 2026")
r.font.size = Pt(9.5)
r.font.color.rgb = RGBColor(0x59, 0x59, 0x59)

# ---- the problem ---------------------------------------------------------
heading("The problem", before=4)
para("Our agent, StochasticGoose (“Goose”), learns by trial and error. After every move it "
     "gets one signal: did the screen change? In most ARC-AGI-3 games the screen changes on almost "
     "every move anyway, because of blinking lights and progress bars. So Goose was told “good job” "
     "nearly all the time and learned very little. It also kept repeating moves it had already made.")

# ---- what we added ------------------------------------------------------
heading("What we added this semester, and why")
bullet("We replayed about 190 saved games and counted how often each kind of signal would have said "
       "“good job.” It cost two days and no new code, so a bad idea would have been dropped cheaply. "
       "The old signal said yes on nearly every move. The new one said yes only occasionally on the games "
       "where Goose was stuck, which gives it something to learn from.",
       "We checked the idea on old recordings first.")
bullet("Goose now learns which spots on the screen flicker constantly and ignores them when comparing "
       "screens. Researchers call this trap the “noisy TV”: a curious agent can stare at random flicker "
       "forever because it always looks new. We reused the rule our scoring tools already use, so the "
       "agent and our reports count the same thing.",
       "We taught Goose to ignore decorations.")
bullet("Goose now gets “good job” only when a move leads to a screen it has not seen before in the "
       "current level. Rewarding rarely visited places is a well-tested way to drive exploration, and one "
       "study did it the same way we do, by turning each screen into a short fingerprint and counting. "
       "Memory resets at each new level, because each level is a new puzzle.",
       "We rewarded new screens, not just changed ones.")
bullet("The second-place entry in the ARC-AGI-3 preview competition removed repeated moves, so it was "
       "worth testing. It hurt us, because some games need the same button pressed many times.",
       "We tried a “don’t repeat yourself” rule, then dropped it.")
bullet("Teammates share this code, and their results cannot change unless they choose to turn it on.",
       "Everything sits behind an off switch.")
bullet("A small round on 6 games picked the winner. A full round on all 25 public games, three tries "
       "each, confirmed it. We wrote down what would count as success before seeing any results.",
       "We tested in two rounds, with the rules written first.")

# ---- results --------------------------------------------------------------
heading("Results")
rows = [("", "Old Goose", "New Goose"),
        ("Levels finished, 75 games played each", "54", "79"),
        ("Games with at least one level finished", "11 of 25", "18 of 25"),
        ("Same game and start, new vs old", "", "23 better, 51 same, 1 worse"),
        ("Speed, moves per second", "140", "139")]
tbl = doc.add_table(rows=len(rows), cols=3)
tbl.style = "Table Grid"
tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
widths = (Inches(3.3), Inches(1.2), Inches(2.1))
for i, row in enumerate(rows):
    for j, val in enumerate(row):
        c = tbl.cell(i, j)
        c.width = widths[j]
        c.text = ""
        p = c.paragraphs[0]
        p.paragraph_format.space_after = Pt(1)
        run = p.add_run(val)
        run.font.size = Pt(10)
        if i == 0:
            run.bold = True
            shade(c, "DCE6F1")
        if j > 0:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
para("Seven games that had never produced a level now do. Two honest limits: the new Goose is not "
     "faster at reaching the first level, and 7 games still produce nothing.", after=2).paragraph_format.space_before = Pt(5)

# ---- next step ------------------------------------------------------------
heading("Next step: give Goose a map and a way back")
para("Goose now wants new screens, but it only thinks one move ahead. If the unexplored part of a level "
     "is several moves away, Goose wanders. It also restarts a lot: on most games it hits a game over and "
     "goes back to the start of the level every 20 to 200 moves, and each time it has to find its way "
     "back by chance.", "What is still wrong.")
para("Goose already fingerprints every screen. We add one thing: remember which move leads from which "
     "screen to which. That is a map, and it has two uses.", "The idea.", after=2)
bullet("When Goose goes a while without finding anything new, it picks the nearest screen on the map "
       "that still has untried moves, walks there by the shortest known route, and explores from there.",
       style="List Number")
bullet("After a game over, it walks straight back to where it last found something new, instead of "
       "starting over.", style="List Number")
para("", "Why this is the right next step.", after=2).paragraph_format.space_before = Pt(6)
bullet("The second- and third-place entries in the ARC-AGI-3 preview both kept a map of screens. "
       "The third-place agent always heads for the shortest route to something untried and, with no "
       "learning at all, solved a median of 30 of 52 levels.",
       "The top entries that did not use language models do this.")
bullet("Go-Explore, published in Nature in 2021, found that agents fail hard exploration games mainly "
       "because they forget how to get back to promising places. Remembering and returning let it beat "
       "famously hard Atari games.",
       "It fixes a well-known cause of failure.")
bullet("The screen fingerprints and memory from this semester’s work are the map’s building blocks.",
       "It builds on what we already made.")
para("In some games two screens look identical but hide a counter, so a remembered route may not work. "
     "Goose checks every step of a route and returns to normal play the moment the game does something "
     "unexpected. Walking back also costs moves, so we measure whether it pays off.", "The risk.")
para("The map lives in its own file behind its own off switch. We compare Goose with and without the "
     "map, first on 6 games and then on all 25, with the success rule written down before any results. "
     "If the map hurts, the switch stays off and the map changes nothing, while its file and results "
     "stay in the project as a documented experiment.", "How we test it, and how it stays removable.")

# ---- sources ---------------------------------------------------------------
heading("Sources", size=11, before=8)
for s in [
    "ARC Prize, “ARC-AGI-3 Preview: 30-Day Learnings,” 2025. arcprize.org/blog/arc-agi-3-preview-30-day-learnings",
    "“Graph-Based Exploration for ARC-AGI-3 Interactive Reasoning Tasks,” 3rd-place entry, 2025. arxiv.org/abs/2512.24156",
    "Ecoffet et al., “First return, then explore,” Nature 590, 2021. nature.com/articles/s41586-020-03157-9",
    "Tang et al., “#Exploration: A Study of Count-Based Exploration for Deep RL,” NeurIPS 2017. arxiv.org/abs/1611.04717",
    "Burda et al., “Large-Scale Study of Curiosity-Driven Learning” (the noisy TV), ICLR 2019. arxiv.org/abs/1808.04355",
]:
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(1)
    run = p.add_run(s)
    run.font.size = Pt(9)

doc.save(OUT)
print("wrote", OUT)
